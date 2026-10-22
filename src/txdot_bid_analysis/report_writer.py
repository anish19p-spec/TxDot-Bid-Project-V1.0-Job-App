"""Create consistent Excel/PDF snapshots from the same verified comparison data."""
from decimal import Decimal
import json
import os
from pathlib import Path
from tempfile import TemporaryDirectory
from xml.sax.saxutils import escape

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import pdfplumber
from reportlab.lib import colors
from reportlab.lib.pagesizes import A3, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.platypus.doctemplate import LayoutError

from .analyzer import analyze_bids
from .errors import ReportError
from .models import AnalysisResult, HistoricalPrice, Project, ReportPaths
from .validator import validate_project

NAVY = "17324D"
GREEN = "DDF0DF"
YELLOW = "FFF1C7"
RED = "F8DFDF"
MONEY = '"$"#,##0.00'


def _money(value: Decimal) -> str:
    precision = max(2, -value.as_tuple().exponent)
    return f"${value:,.{precision}f}"


def _names(project: Project, ids: tuple[str, ...]) -> str:
    names = {bidder.bidder_id: bidder.company_name for bidder in project.bidders}
    prefix = "Tied: " if len(ids) > 1 else ""
    return prefix + "; ".join(names[key] for key in ids) if ids else "-"


def _comparison_rows(project: Project, analysis: AnalysisResult) -> list[list]:
    records = {(r.bidder_id, r.item_id): r for r in project.records}
    comparisons = {c.item_id: c for c in analysis.item_comparisons}
    rows = []
    for item in project.items:
        row = [item.item_number, item.quantity, item.unit, item.description]
        for bidder in project.bidders:
            record = records[bidder.bidder_id, item.item_id]
            row.extend([record.unit_price, record.extended_amount])
        comp = comparisons[item.item_id]
        row.extend([_names(project, comp.lowest_bidder_ids), _names(project, comp.middle_bidder_ids),
                    _names(project, comp.highest_bidder_ids)])
        rows.append(row)
    return rows


def _set(cell, value) -> None:
    cell.value = value
    if isinstance(value, str):
        # Untrusted PDF names/descriptions must remain literal Excel text, never formulas.
        cell.data_type = "s"
    cell.alignment = Alignment(vertical="center", wrap_text=True)


def _excel(project: Project, analysis: AnalysisResult, rows: list[list], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Bid comparison"
    n = 4 + 2 * len(project.bidders) + 3
    for row, text in [(1, f"{project.airport_name} - Bid comparison"),
                      (2, f"{project.project_id} | Bid date: {project.bid_date.isoformat()}"),
                      (3, "Lowest submitted bidder: " + _names(project, analysis.lowest_submitted_bidder_ids)),
                      (4, "Item colors: green lowest, yellow middle, red highest. Ties retained. No award status inferred."),
                      (5, "Static snapshot: regenerate reports after changes; editing cells does not recalculate rankings.")]:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n)
        _set(ws.cell(row, 1), text)
        ws.row_dimensions[row].height = 27 if row == 1 else 24
    ws.cell(1, 1).font = Font(size=20, bold=True, color=NAVY)
    ws.cell(3, 1).font = Font(size=12, bold=True, color="206538")
    for col, title in enumerate(["Item", "Quantity", "Unit", "Description"], 1):
        ws.merge_cells(start_row=7, start_column=col, end_row=8, end_column=col)
        _set(ws.cell(7, col), title)
    for i, bidder in enumerate(project.bidders):
        col = 5 + 2 * i
        ws.merge_cells(start_row=7, start_column=col, end_row=7, end_column=col+1)
        _set(ws.cell(7, col), bidder.company_name)
        _set(ws.cell(8, col), "Unit price")
        _set(ws.cell(8, col+1), "Submitted amount")
    for col, label in enumerate(["Lowest (unit price)", "Middle / intermediate", "Highest (unit price)"], n-2):
        ws.merge_cells(start_row=7, start_column=col, end_row=8, end_column=col)
        _set(ws.cell(7, col), label)
    for cells in ws.iter_rows(min_row=7, max_row=8, max_col=n):
        for cell in cells:
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(bold=True, color="FFFFFF", size=11)
    ws.row_dimensions[7].height = 40
    ws.row_dimensions[8].height = 28
    for r, values in enumerate(rows, 9):
        for col, value in enumerate(values, 1):
            cell = ws.cell(r, col)
            _set(cell, value)
            cell.font = Font(name="Calibri", size=11)
            if isinstance(value, Decimal):
                if col == 2:
                    cell.number_format = '#,##0' if value == value.to_integral_value() else '#,##0.####'
                else:
                    cell.number_format = '"$"#,##0.' + '0' * max(2, -value.as_tuple().exponent)
        ws.cell(r, 2).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(r, 3).alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[r].height = 64
        comp = analysis.item_comparisons[r-9]
        for i, bidder in enumerate(project.bidders):
            fill = GREEN if bidder.bidder_id in comp.lowest_bidder_ids else RED if bidder.bidder_id in comp.highest_bidder_ids else YELLOW
            for col in (5 + 2*i, 6 + 2*i):
                ws.cell(r, col).fill = PatternFill("solid", fgColor=fill)
    total_row = 9 + len(rows)
    _set(ws.cell(total_row, 4), "Submitted base bid")
    ranks = {r.bidder_id: r.rank for r in analysis.rankings}
    for i, bidder in enumerate(project.bidders):
        col = 6 + 2*i
        _set(ws.cell(total_row, col), bidder.submitted_total)
        ws.cell(total_row, col).number_format = MONEY
        ws.cell(total_row, col).font = Font(bold=True)
        if bidder.bidder_id in analysis.lowest_submitted_bidder_ids:
            ws.cell(total_row, col).fill = PatternFill("solid", fgColor=GREEN)
        _set(ws.cell(total_row+1, col), f"Overall rank {ranks[bidder.bidder_id]}")
    for r in (total_row, total_row+1):
        ws.row_dimensions[r].height = 26
    source_row = total_row + 3
    sources = ["Source PDF: " + project.source_url, "Listing (airport/date): " + project.listing_url,
               "SHA-256: " + project.document.sha256,
               "Physical table pages by bidder: " + "; ".join(f"{b.company_name}: " + ", ".join(str(p) for p in sorted({r.source.page for r in project.records if r.bidder_id == b.bidder_id})) for b in project.bidders)]
    sources.extend("Warning: " + warning.message for warning in analysis.warnings)
    for offset, value in enumerate(sources):
        row = source_row + offset
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n)
        _set(ws.cell(row, 1), value)
        ws.row_dimensions[row].height = 30
    widths = [14, 13, 9, 49] + [19, 21] * len(project.bidders) + [33, 33, 33]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "E9"
    ws.sheet_view.showGridLines = False
    ws.print_title_rows = "1:8"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_area = f"A1:{get_column_letter(n)}{source_row+len(sources)-1}"
    wb.save(path)
    wb.close()


def _pdf(project: Project, analysis: AnalysisResult, rows: list[list], path: Path) -> None:
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="Cell", fontName="Helvetica", fontSize=9, leading=12))
    styles.add(ParagraphStyle(name="SmallCell", fontName="Helvetica", fontSize=8, leading=10))
    def p(value, small=False):
        if isinstance(value, Decimal):
            value = _money(value)
        return Paragraph(escape(str(value)), styles["SmallCell" if small else "Cell"])
    story = [Paragraph(escape(project.airport_name), styles["Title"]),
             p(f"{project.project_id} | Bid date {project.bid_date.isoformat()}"), Spacer(1, 10),
             Paragraph("Lowest submitted bidder: " + escape(_names(project, analysis.lowest_submitted_bidder_ids)), styles["Heading2"]),
             p("Base bids. Item ranking uses unit price. Ties retained. No award status inferred."), Spacer(1, 15)]
    # Split very wide projects into readable bidder groups, retaining shared item columns.
    for start in range(0, len(project.bidders), 3):
        group = project.bidders[start:start+3]
        if start:
            story.extend([PageBreak(), Paragraph("Bid comparison (continued)", styles["Heading1"])])
        header = [p(x) for x in ["Item", "Qty", "Unit", "Description"]]
        for bidder in group:
            header.extend([p(bidder.company_name), p("")])
        subheader = [p("") for _ in range(4)] + [p(value) for _ in group for value in ("Unit price", "Submitted amount")]
        data = [header, subheader]
        for row in rows:
            data.append([p(row[0]), p(f"{row[1]:,}"), p(row[2]), p(row[3])] +
                        [p(value) for i in range(start, start+len(group)) for value in row[4+2*i:6+2*i]])
        data.append([p(""), p(""), p(""), p("Submitted base bid")] +
                    [p(value) for bidder in group for value in ("", bidder.submitted_total)])
        widths = [70, 55, 42, 245] + [105, 115] * len(group)
        table = Table(data, colWidths=widths, repeatRows=2, hAlign="LEFT")
        commands = [("BACKGROUND", (0,0), (-1,1), colors.HexColor("#E6EDF4")),
                    ("GRID", (0,0), (-1,-1), .35, colors.HexColor("#BBC7D2")),
                    ("VALIGN", (0,0), (-1,-1), "TOP"),
                    ("TOPPADDING", (0,0), (-1,-1), 7), ("BOTTOMPADDING", (0,0), (-1,-1), 7)]
        commands.extend(("SPAN", (col,0), (col,1)) for col in range(4))
        commands.extend(("SPAN", (4+2*i,0), (5+2*i,0)) for i in range(len(group)))
        for r, comparison in enumerate(analysis.item_comparisons, 2):
            for i, bidder in enumerate(group):
                fill = GREEN if bidder.bidder_id in comparison.lowest_bidder_ids else RED if bidder.bidder_id in comparison.highest_bidder_ids else YELLOW
                commands.append(("BACKGROUND", (4+2*i,r), (5+2*i,r), colors.HexColor("#"+fill)))
        table.setStyle(TableStyle(commands))
        story.append(table)
    story.extend([PageBreak(), Paragraph("Item price positions", styles["Title"]),
                  p("Companies are ordered independently for each item. A dash means no bidder is strictly between the lowest and highest prices."), Spacer(1, 12)])
    positions = [[p(x) for x in ["Item / description", "Lowest", "Middle / intermediate", "Highest"]]]
    for row in rows:
        positions.append([p(f"{row[0]} - {row[3]}")] + [p(v) for v in row[-3:]])
    table = Table(positions, colWidths=[350, 245, 245, 245], repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#E6EDF4")),
        ("GRID",(0,0),(-1,-1),.35,colors.HexColor("#BBC7D2")), ("VALIGN",(0,0),(-1,-1),"TOP"),
        ("TOPPADDING",(0,0),(-1,-1),8), ("BOTTOMPADDING",(0,0),(-1,-1),8)]))
    story.extend([table, Spacer(1, 16), p("Source PDF: " + project.source_url, True),
                  p("Listing (airport/date): " + project.listing_url, True),
                  p("SHA-256: " + project.document.sha256, True),
                  p("Static snapshot. Regenerate both reports after changes.", True)])
    for issue in analysis.warnings:
        story.append(p("Warning: " + issue.message))
    def footer(canvas, doc):
        canvas.setFont("Helvetica", 8)
        canvas.drawString(36, 20, f"{project.project_id} | Page {doc.page}")
    SimpleDocTemplate(str(path), pagesize=landscape(A3), leftMargin=36, rightMargin=36,
                      topMargin=30, bottomMargin=35).build(story, onFirstPage=footer, onLaterPages=footer)


def write_reports(project: Project, analysis: AnalysisResult, output_dir: Path) -> ReportPaths:
    """Recompute and verify analysis before staging both static report snapshots."""
    current = analyze_bids(project, validate_project(project))
    if analysis != current:
        raise ReportError("Analysis does not match the current project; regenerate rankings first.")
    try:
        output_dir.mkdir(parents=True, exist_ok=True)
        rows = _comparison_rows(project, analysis)
        paths = ReportPaths(output_dir / "bid_comparison.xlsx", output_dir / "bid_comparison.pdf")
        with TemporaryDirectory(dir=output_dir) as folder:
            excel, pdf = Path(folder) / paths.excel.name, Path(folder) / paths.pdf.name
            _excel(project, analysis, rows, excel)
            _pdf(project, analysis, rows, pdf)
            wb = load_workbook(excel, data_only=True)
            try:
                ws = wb.active
                for r, row in enumerate(rows, 9):
                    for c, expected in enumerate(row, 1):
                        actual = ws.cell(r,c).value
                        differs = Decimal(str(actual)) != expected if isinstance(expected, Decimal) else actual != expected
                        if differs:
                            raise ReportError("Saved workbook differs from validated report values.")
            finally:
                wb.close()
            with pdfplumber.open(pdf) as saved:
                text = "\n".join(page.extract_text() or "" for page in saved.pages)
                if project.project_id not in text or any(f"{b.submitted_total:,.2f}" not in text for b in project.bidders):
                    raise ReportError("Saved PDF is missing required project or submitted totals.")
            # The manifest is the completion marker; two file replacements are not one transaction.
            marker = output_dir / "report_manifest.json"
            marker.unlink(missing_ok=True)
            os.replace(excel, paths.excel)
            os.replace(pdf, paths.pdf)
            manifest = Path(folder) / marker.name
            manifest.write_text(json.dumps({"project_id": project.project_id,
                "source_sha256": project.document.sha256, "excel": paths.excel.name,
                "pdf": paths.pdf.name, "status": "reports_complete_master_not_updated"}, indent=2))
            os.replace(manifest, marker)
        return paths
    except ReportError:
        raise
    except (OSError, ValueError, ArithmeticError, LayoutError) as exc:
        raise ReportError("Could not write comparison reports; close open workbooks and check folder permissions.") from exc



def write_estimate_workbook(prices: tuple[HistoricalPrice, ...], path: Path) -> Path:
    """Build live Excel formulas over a read-only snapshot of historical observations."""
    from datetime import datetime
    from zoneinfo import ZoneInfo
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.workbook.defined_name import DefinedName
    from .estimator import prepare_observations
    from .master_sheet import _same_cells

    observations = prepare_observations(prices)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Estimate'
    history = wb.create_sheet('History')
    last = len(observations) + 1
    keys = sorted({o.item_key for o in observations})
    regions = ['All regions'] + sorted({o.sources[0].region for o in observations} | {'Unknown'})
    labels = {
        1: 'Historical item cost estimate', 2: 'Edit yellow cells. Prices come from the saved master.',
        3: 'Item / description / unit', 4: 'Planned quantity', 5: 'Estimate date',
        6: 'Rolling months', 7: 'Quantity tolerance (optional)', 8: 'Region (TxDOT district)', 9: 'Input status',
        10: 'Earliest eligible date', 11: 'Matching projects', 12: 'Median unit price',
        13: 'Estimated item cost', 14: 'Minimum project unit price', 15: 'Maximum project unit price',
        16: 'Result basis', 17: 'History rows with unknown region', 18: 'Oldest matching bid', 19: 'Newest matching bid',
        21: 'How to use', 22: 'Interpretation', 23: 'Comparability', 24: 'Refresh history',
    }
    for row, text in labels.items():
        _set(ws.cell(row, 1), text)
        ws.row_dimensions[row].height = 30
    for row in (1, 2):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    ws['A1'].font = Font(size=16, bold=True, color=NAVY)
    inputs = {3: keys[0], 4: observations[0].quantity, 5: datetime.now(ZoneInfo('America/Chicago')).date(),
              6: 12, 7: None, 8: 'All regions'}
    for row, value in inputs.items():
        _set(ws.cell(row, 2), value)
        ws.cell(row, 2).fill = PatternFill('solid', fgColor=YELLOW)
    ws['B5'].number_format = 'mmmm d, yyyy'
    ws['B7'].number_format = '0%'
    # Guard pasted values too: Excel input validation alone can be bypassed by pasting.
    ws['B9'] = ('=IF(AND(COUNTIF(History!$L$2:$L$'+str(len(keys)+1)+',$B$3)=1,'
                f'COUNTIF(History!$O$2:$O${len(regions)+1},$B$8)=1,'
                'ISNUMBER(B4),B4>0,ISNUMBER(B5),B5=INT(B5),B5>=DATE(1901,1,1),B5<=DATE(9998,12,31),'
                'OR(B6=3,B6=6,B6=12),OR(B7="",AND(ISNUMBER(B7),B7>=0,B7<=1))),"Ready","Check inputs")')
    ws['B9'] = '=IFERROR(' + ws['B9'].value[1:] + ',"Check inputs")'
    ws['B10'] = '=IF(B9="Ready",EDATE(B5,-B6),"")'
    match_prices = f'History!$J$2:$J${last}'
    match_dates = f'History!$K$2:$K${last}'
    ws['B11'] = f'=IF(B9="Ready",COUNT({match_prices}),0)'
    for row, expression in {12:f'MEDIAN({match_prices})', 13:'ROUND(B12*B4,2)',
                            14:f'MIN({match_prices})', 15:f'MAX({match_prices})',
                            18:f'MIN({match_dates})', 19:f'MAX({match_dates})'}.items():
        ws.cell(row, 2).value = f'=IF(B11=0,"",{expression})'
    ws['B17'] = f'=COUNTIF(History!$M$2:$M${last},"Unknown")'
    ws['B16'] = '=IF(B9<>"Ready","Check inputs",IF(B11=0,"No matching data",IF(B11=1,"Single-project reference","Historical median")))'
    for row in (10, 18, 19):
        ws.cell(row, 2).number_format = 'mmmm d, yyyy'
    for row in (12, 13, 14, 15):
        ws.cell(row, 2).number_format = '"$"#,##0.00'
    ws['B13'].fill = PatternFill('solid', fgColor=GREEN)
    notes = {
        21: 'Choose an item, enter your quantity and date, then select 3, 6 or 12 rolling months. '
            'Leave tolerance blank for any historical quantity, or enter 25% for +/-25%. Initial quantity is from history; replace it with your planned quantity.',
        22: 'Historical prices are not inflation-adjusted current quotes. One project is a reference, not a market forecast. '
            'Min/max are observed project prices, not confidence limits. No regression is used.',
        23: 'Exact normalized item, description and unit matches only. Review specifications before relying on comparisons. '
            'Tied low bidders contribute one median price per project. Select a TxDOT district in B8. Unknown locations match only All regions or Unknown; they are excluded from district estimates.',
        24: 'Inputs recalculate in Excel. History is a snapshot; rerun --estimate after updating the master. '
            'Regeneration replaces this estimate file and resets its inputs. Save a separate copy to retain an estimate.',
    }
    for row, value in notes.items():
        _set(ws.cell(row, 2), value)
        ws.row_dimensions[row].height = 85
    ws.column_dimensions['A'].width = 35
    ws.column_dimensions['B'].width = 95
    ws.row_dimensions[3].height = 48
    for row in range(3, 25):
        ws.cell(row, 2).alignment = Alignment(vertical='center', wrap_text=True)
    ws.sheet_view.showGridLines = False

    headers = ('Item / description / unit', 'Project ID', 'Bid date', 'Quantity', 'Project unit price',
               'PDF URL', 'SHA-256', 'Source pages', 'Low bidder prices (ID = price)',
               'Matching unit price', 'Matching bid date', 'Available items', 'TxDOT district', 'County', 'Available regions', 'Geography sources', 'Airport')
    for col, label in enumerate(headers, 1):
        _set(history.cell(1, col), label)
        history.cell(1, col).fill = PatternFill('solid', fgColor=NAVY)
        history.cell(1, col).font = Font(bold=True, color='FFFFFF')
        history.column_dimensions[get_column_letter(col)].width = {1:65, 6:45, 7:35, 9:50, 12:65, 16:65, 17:35}.get(col, 22)
    history.row_dimensions[1].height = 32
    for row, observation in enumerate(observations, 2):
        first = observation.sources[0]
        values = (observation.item_key, observation.project_id, observation.bid_date,
                  observation.quantity, observation.unit_price, first.source_url, first.source_hash,
                  ', '.join(str(s.source_page) for s in observation.sources),
                  '; '.join(f'{s.bidder_id} = {s.unit_price}' for s in observation.sources))
        for col, value in enumerate(values, 1):
            _set(history.cell(row, col), value)
        history.cell(row, 3).number_format = 'yyyy-mm-dd'
        history.cell(row, 5).number_format = '"$"#,##0.00'
        history.cell(row, 10).number_format = '"$"#,##0.00'
        history.cell(row, 11).number_format = 'yyyy-mm-dd'
        for column, value in ((13, first.region), (14, first.county), (16, first.geography_sources), (17, first.airport_name)):
            _set(history.cell(row, column), value or None)
        eligible = (f'AND(A{row}=Estimate!$B$3,C{row}>=Estimate!$B$10,C{row}<=Estimate!$B$5,'
                    f'OR(Estimate!$B$8="All regions",M{row}=Estimate!$B$8),'
                    f'OR(Estimate!$B$7="",AND(D{row}>=Estimate!$B$4*(1-Estimate!$B$7),'
                    f'D{row}<=Estimate!$B$4*(1+Estimate!$B$7))))')
        history.cell(row, 10).value = f'=IF(Estimate!$B$9<>"Ready","",IF({eligible},E{row},""))'
        history.cell(row, 11).value = f'=IF(ISNUMBER(J{row}),C{row},"")'
        history.row_dimensions[row].height = 95
    for row, key in enumerate(keys, 2):
        _set(history.cell(row, 12), key)
    for row, region in enumerate(regions, 2):
        _set(history.cell(row, 15), region)
    wb.defined_names.add(DefinedName('EstimateRegions', attr_text=f"'History'!$O$2:$O${len(regions)+1}"))
    wb.defined_names.add(DefinedName('EstimateItems', attr_text=f"'History'!$L$2:$L${len(keys)+1}"))
    for cell, validation in [
        ('B8', DataValidation(type='list', formula1='EstimateRegions')),
        ('B3', DataValidation(type='list', formula1='EstimateItems')),
        ('B4', DataValidation(type='decimal', operator='greaterThan', formula1='0')),
        ('B5', DataValidation(type='date', operator='between', formula1='DATE(1901,1,1)', formula2='DATE(9998,12,31)')),
        ('B6', DataValidation(type='list', formula1='"3,6,12"')),
        ('B7', DataValidation(type='decimal', operator='between', formula1='0', formula2='1', allow_blank=True)),
    ]:
        validation.showErrorMessage = True
        validation.errorTitle = 'Check estimate input'
        validation.error = 'Use a value allowed by this field.'
        validation.errorStyle = 'stop'
        ws.add_data_validation(validation)
        validation.add(ws[cell])
    history.auto_filter.ref = f'A1:K{last}'
    history.freeze_panes = 'B2'
    history.sheet_view.showGridLines = False
    wb.calculation.calcMode = 'auto'
    wb.calculation.fullCalcOnLoad = True
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        with TemporaryDirectory(dir=path.parent) as folder:
            temporary = Path(folder)/path.name
            wb.save(temporary)
            check = load_workbook(temporary)
            try:
                if not _same_cells(check, wb):
                    raise ReportError('Estimate workbook failed round-trip verification.')
            finally:
                check.close()
            os.replace(temporary, path)
    except ReportError:
        raise
    except Exception:
        raise ReportError('Could not save estimate. Close Excel and check output permissions.') from None
    finally:
        wb.close()
    return path
