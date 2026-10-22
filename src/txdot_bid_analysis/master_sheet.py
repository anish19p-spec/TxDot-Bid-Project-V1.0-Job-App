"""Safe, repeatable storage of overall low bidders in a horizontal master workbook."""
from datetime import date, datetime
from decimal import Decimal
import hashlib
import os
from pathlib import Path
import shutil
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .analyzer import analyze_bids
from .errors import MasterUpdateError
from .models import AnalysisResult, HistoricalPrice, MasterUpdateResult, Project
from .validator import validate_project

HEADERS = ("Project ID", "Airport", "Bid date", "Bidder ID", "Company", "Item ID", "Item",
           "Description", "Unit", "Quantity", "Unit price", "PDF URL", "Listing URL", "SHA-256",
           "PDF page", "Submitted bid total")
LEGACY_SCHEMA = "TxDOT low-bid master v1"
SCHEMA = "TxDOT low-bid master v2"


def _text(cell, value) -> None:
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"
    cell.alignment = Alignment(vertical="center", wrap_text=True)


def _numeric(value) -> Decimal:
    result = Decimal(str(value))
    if not result.is_finite() or result < 0:
        raise MasterUpdateError("Master contains invalid numeric values.")
    return result


def _read_records(wb) -> list[tuple]:
    if wb.sheetnames != ["Master unit prices", "Records"] or wb.worksheets[0]['A1'].value not in (LEGACY_SCHEMA, SCHEMA):
        raise MasterUpdateError("Unrecognized master workbook; original file was not changed.")
    ws = wb["Records"]
    if tuple(c.value for c in ws[1]) != HEADERS:
        raise MasterUpdateError("Master Records headers do not match the supported schema.")
    records = []
    keys = set()
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = list(values)
        if any(value is None for value in row):
            raise MasterUpdateError("Master contains an incomplete record.")
        if isinstance(row[2], datetime):
            row[2] = row[2].date()
        if not isinstance(row[2], date):
            raise MasterUpdateError("Master bid dates must be Excel dates.")
        for index in (9, 10, 15):
            row[index] = _numeric(row[index])
        key = row[0], row[3], row[5]
        if key in keys:
            raise MasterUpdateError("Master contains duplicate project/bidder/item records.")
        keys.add(key)
        records.append(tuple(row))
    if not records:
        raise MasterUpdateError("Existing master has no historical records.")
    return records


def _build(records: list[tuple], *, legacy: bool = False) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Master unit prices"
    _text(ws['A1'], LEGACY_SCHEMA if legacy else SCHEMA)
    wb.calculation.calcMode = 'auto'
    wb.calculation.fullCalcOnLoad = True
    ws['A1'].font = Font(size=18, bold=True, color="17324D")
    ws.merge_cells('A1:C1')
    _text(ws['A3'], "Overall lowest submitted bidders only; ties retained.")
    ws.merge_cells('A3:C3')
    _text(ws['A4'], "Blank = item not present in that project. Dates come from the listing.")
    ws.merge_cells('A4:C4')
    for col, title in enumerate(("Item", "Description", "Unit"), 1):
        _text(ws.cell(6, col), title)
    # Match only exact item number/description/unit combinations; never merge unlike work by name.
    item_rows = {}
    blocks = {}
    for record in records:
        key = record[6].strip().upper(), ' '.join(record[7].split()).casefold(), record[8].strip().upper()
        if key not in item_rows:
            row = 7 + len(item_rows)
            item_rows[key] = row
            for col, value in enumerate((record[6], record[7], record[8]), 1):
                _text(ws.cell(row, col), value)
            ws.row_dimensions[row].height = 42
        block = record[0], record[3]
        if block not in blocks:
            col = 4 + 2 * len(blocks)
            blocks[block] = col
            for row, value in [(2, f"{record[1]} ({record[2]:%B %Y})"), (3, record[4]),
                               (4, record[2]), (5, record[0])]:
                if legacy or row != 4:
                    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col+1)
                _text(ws.cell(row, col), value)
            ws.cell(4, col).number_format = 'mmmm d, yyyy'
            if not legacy:
                ref = f'{get_column_letter(col)}4'
                age = ws.cell(4, col+1)
                age.value = (f'=IF({ref}=TODAY(),"Today",IF({ref}<TODAY(),'
                             f'TODAY()-{ref}&" days ago","In "&({ref}-TODAY())&" days"))')
                age.alignment = Alignment(vertical="center", wrap_text=True)
            _text(ws.cell(6, col), "Unit Price")
            _text(ws.cell(6, col+1), "Qty")
        col = blocks[block]
        row = item_rows[key]
        if ws.cell(row, col).value is not None:
            raise MasterUpdateError("Two records map to the same item row in one project block.")
        _text(ws.cell(row, col), record[10])
        ws.cell(row, col).number_format = '"$"#,##0.' + '0' * max(2, -record[10].as_tuple().exponent)
        ws.cell(row, col).fill = PatternFill('solid', fgColor='DDF3F7')
        _text(ws.cell(row, col+1), record[9])
        ws.cell(row, col+1).number_format = '#,##0' if record[9] == record[9].to_integral_value() else '#,##0.####'
    for col in range(1, 4 + 2*len(blocks)):
        cell = ws.cell(6, col)
        cell.fill = PatternFill('solid', fgColor='17324D')
        cell.font = Font(bold=True, color='FFFFFF')
        ws.column_dimensions[get_column_letter(col)].width = {1:16, 2:65, 3:10}.get(col, 24)
    for row, height in [(2,40),(3,35),(4,30),(5,25),(6,30)]:
        ws.row_dimensions[row].height = height
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'D7'
    ws.print_title_rows = '1:6'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    source = wb.create_sheet("Records")
    for col, title in enumerate(HEADERS, 1):
        _text(source.cell(1,col), title)
        source.cell(1,col).fill = PatternFill('solid', fgColor='17324D')
        source.cell(1,col).font = Font(bold=True,color='FFFFFF')
        source.column_dimensions[get_column_letter(col)].width = 22 if col not in (2,5,8,12,13,14) else 44
    source.row_dimensions[1].height = 30
    for row, record in enumerate(records, 2):
        source.row_dimensions[row].height = 65
        for col, value in enumerate(record, 1):
            _text(source.cell(row,col), value)
        source.cell(row,3).number_format = 'yyyy-mm-dd'
        for col in (11,16):
            source.cell(row,col).number_format = '"$"#,##0.00'
    source.freeze_panes = 'A2'
    source.auto_filter.ref = source.dimensions
    return wb


def _same_cells(actual, expected) -> bool:
    if actual.sheetnames != expected.sheetnames:
        return False
    for left, right in zip(actual, expected):
        if (left.max_row, left.max_column) != (right.max_row, right.max_column):
            return False
        for lrow, rrow in zip(left, right):
            for lc, rc in zip(lrow, rrow):
                a, b = lc.value, rc.value
                if isinstance(a, datetime): a = a.date()
                if isinstance(b, Decimal):
                    if a is None or Decimal(str(a)) != b: return False
                elif a != b:
                    return False
    return True


def update_master(project: Project, analysis: AnalysisResult, path: Path) -> MasterUpdateResult:
    """Add only overall low bidders, preserve dates/provenance and safely deduplicate retries."""
    if analyze_bids(project, validate_project(project)) != analysis:
        raise MasterUpdateError("Analysis is stale; recalculate the lowest submitted bidder first.")
    lowest_bidder_ids = set(analysis.lowest_submitted_bidder_ids)
    items = {item.item_id:item for item in project.items}
    bidders = {bidder.bidder_id:bidder for bidder in project.bidders}
    incoming = []
    for record in project.records:
        if record.bidder_id not in lowest_bidder_ids:
            continue
        lowest_bidder = bidders[record.bidder_id]
        item = items[record.item_id]
        incoming.append((project.project_id, project.airport_name, project.bid_date,
            lowest_bidder.bidder_id, lowest_bidder.company_name, item.item_id, item.item_number,
            item.description, item.unit, item.quantity, record.unit_price, project.source_url,
            project.listing_url, project.document.sha256, record.source.page, lowest_bidder.submitted_total))
    lock = path.with_suffix(path.suffix + '.lock')
    acquired = False
    try:
        if hashlib.sha256(project.document.path.read_bytes()).hexdigest() != project.document.sha256:
            raise MasterUpdateError("Original PDF content no longer matches its source hash.")
        path.parent.mkdir(parents=True, exist_ok=True)
        fd = os.open(lock, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
        os.close(fd)
        acquired = True
        records = []
        upgrade = False
        if path.exists():
            wb = load_workbook(path)
            try:
                records = _read_records(wb)
                upgrade = wb.worksheets[0]['A1'].value == LEGACY_SCHEMA
                rebuilt = _build(records, legacy=upgrade)
                try:
                    if not _same_cells(wb, rebuilt):
                        raise MasterUpdateError("Master display and Records disagree; original was not changed.")
                finally:
                    rebuilt.close()
            finally:
                wb.close()
        existing = [row for row in records if row[0] == project.project_id]
        if existing:
            if set(existing) == set(incoming):
                if not upgrade:
                    return MasterUpdateResult(path, 'unchanged', 0)
            else:
                raise MasterUpdateError("Project already exists with changed source or data; review the revision before updating.")
        added = [] if existing else incoming
        wb = _build(records + added)
        try:
            with TemporaryDirectory(dir=path.parent) as folder:
                temporary = Path(folder)/path.name
                wb.save(temporary)
                check = load_workbook(temporary)
                try:
                    _read_records(check)
                    if not _same_cells(check, wb):
                        raise MasterUpdateError("Saved master values failed round-trip verification.")
                finally:
                    check.close()
                if path.exists():
                    backup = Path(folder)/'backup.xlsx'
                    shutil.copy2(path, backup)
                    os.replace(backup, path.with_suffix('.bak.xlsx'))
                os.replace(temporary, path)
        finally:
            wb.close()
        return MasterUpdateResult(path, 'upgraded' if existing else 'inserted', len(added))
    except MasterUpdateError:
        raise
    except FileExistsError:
        raise MasterUpdateError("Another master update holds the lock; retry after it finishes.") from None
    except Exception:
        raise MasterUpdateError("Could not update master. Close Excel and check permissions/workbook integrity; original retained.") from None
    finally:
        if acquired:
            lock.unlink(missing_ok=True)



def read_historical_prices(path: Path) -> tuple[HistoricalPrice, ...]:
    """Read validated master values without changing the workbook."""
    try:
        wb = load_workbook(path)
        try:
            records = _read_records(wb)
            expected = _build(records, legacy=wb.worksheets[0]['A1'].value == LEGACY_SCHEMA)
            try:
                if not _same_cells(wb, expected):
                    raise MasterUpdateError("Master display and Records disagree.")
            finally:
                expected.close()
            return tuple(HistoricalPrice(r[0], r[3], r[2], r[6], r[7], r[8], r[9], r[10],
                                         r[11], r[13], r[14], airport_name=r[1]) for r in records)
        finally:
            wb.close()
    except MasterUpdateError:
        raise
    except Exception:
        raise MasterUpdateError("Could not read master; check its path and workbook integrity.") from None
